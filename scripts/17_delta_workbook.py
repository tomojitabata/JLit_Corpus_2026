#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
17_delta_workbook.py
====================
**Burrows's Delta を Excel で手計算するためのブック**を作る（Step 4 §2）。

Delta の計算は三つの手順だけからなる。

1. 最頻語 N 語の相対頻度を求める（ここでは 1 万語あたり）
2. 語ごとに z スコアにする（平均と標準偏差は**既知の作品だけ**から求める）
3. 問題のテクストと各候補との z の差の絶対値を，語について平均する

ノートブックでは 1 行で済むこの計算を，**1 手順＝1 シート**に展開し，
すべてのセルを数式で書く。度数（青字）を書き換えれば全体が再計算される。

シート
------
    説明            手順・凡例・どこを触ってよいか
    1_度数          最頻語 × 作品の生の度数と総語数（入力値・青字）
    2_相対頻度      1 万語あたり  =度数/総語数*10000
    3_平均と標準偏差 語ごとの AVERAGE・STDEVP（既知の作品だけ）
    4_zスコア       =(相対頻度−平均)/標準偏差
    5_z差           |作品X − 候補| を語ごとに（作品ごと・作家プロファイルごと）
    6_Delta         z差の平均＝Delta と順位。距離をいちばん大きくしている語
    7_語数Nを変える  上位 N 語だけで Delta を計算し直す（N は黄色のセル）

**最頻語と平均・標準偏差は既知の作品だけから決める。** 問題のテクストが
物差しそのものを動かさないようにするためである（Burrows 2002 の手順）。
標準偏差は母標準偏差（STDEVP，numpy の既定 ddof=0）で，
07_descriptive_stats.py と揃えてある。

使い方
------
    python3 scripts/17_delta_workbook.py \\
        --tokens data/tokens/tokens_lemma \\
        --out my_work/results/Step4_delta_manual.xlsx

    # 作品を変える（「作家:題」をカンマで区切る。題は title_aozora）
    python3 scripts/17_delta_workbook.py --mfw 30 \\
        --questioned 夏目漱石:三四郎 \\
        --known 夏目漱石:こころ,夏目漱石:坊っちゃん,森鴎外:青年,森鴎外:ヰタ・セクスアリス,島崎藤村:破戒,島崎藤村:家
    # （既定は 漱石・鴎外・乱歩『少年探偵団』『宇宙怪人』・藤村『千曲川のスケッチ』『新生』の 4 作家）

ブックと同じ場所に ``*_expected.json`` も書く。Python で計算した
正解値（Delta・順位）で，ノートブックの照合に使う。
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
from collections import Counter

import numpy as np

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DEFAULT_QUESTIONED = '夏目漱石:三四郎'
# 既知の作品（4 作家 × 2 点）。漱石・鴎外は『三四郎』と同時期（1906–1914），
# 藤村は同時期だが小説『新生』と写生文『千曲川のスケッチ』，乱歩は時代も
# 読者層（少年もの）も離れた対照。乱歩の『灰色の巨人』は v1 で本文が
# 『魔法博士』と取り違えられていた作品なので，教材には使わない。
DEFAULT_KNOWN = ('夏目漱石:こころ,夏目漱石:坊っちゃん,'
                 '森鴎外:青年,森鴎外:ヰタ・セクスアリス,'
                 '江戸川乱歩:少年探偵団,江戸川乱歩:宇宙怪人,'
                 '島崎藤村:千曲川のスケッチ,島崎藤村:新生')


def default_meta() -> str:
    """自分で作った v3（*_local.csv）> 配布版 v3 > v2。"""
    base = os.path.join(ROOT, 'metadata')
    for name in ('corpus_metadata_v3_local.csv', 'corpus_metadata_v3.csv',
                 'corpus_metadata_v2.csv'):
        p = os.path.join(base, name)
        if os.path.exists(p):
            return p
    return os.path.join(base, 'corpus_metadata_v3.csv')


def disp(w: str) -> str:
    """表示用の語形。UniDic は外来語の語彙素に原綴を付ける（テーブル-table）ので，
    片仮名だけにする。原綴はセルのコメントに残す（解析の誤りを確かめるため）。"""
    import re
    return re.sub(r"-[A-Za-z][A-Za-z .'-]*$", '', w)


def resolve(spec: str, rows: list[dict], tokdir: str) -> dict:
    """「作家:題」からメタデータの行とトークンファイルを引く。"""
    if ':' not in spec:
        sys.exit(f'作品の指定は「作家:題」の形にすること: {spec}')
    author, title = (s.strip() for s in spec.split(':', 1))
    hits = []
    for r in rows:
        if r.get('author_ja') != author or r.get('title_aozora') != title:
            continue
        pid = str(r.get('aozora_person_id') or '').strip()
        wid = str(r.get('aozora_work_id') or '').strip()
        if not (pid.isdigit() and wid.isdigit()):
            continue
        stem = f'{pid.zfill(6)}_{wid.zfill(6)}'
        path = os.path.join(tokdir, stem + '.txt')
        if os.path.exists(path):
            hits.append({'author': author, 'title': title, 'stem': stem,
                         'path': path, 'year': r.get('year_first', ''),
                         'complete': r.get('completeness', '')})
    if not hits:
        sys.exit(f'見つからない: {spec}\n'
                 '  author_ja と title_aozora の綴りをメタデータで確かめること。'
                 '分冊を結合した作品は結合後の題で指定する。')
    return hits[0]


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--tokens', default=os.path.join(ROOT, 'data', 'tokens', 'tokens_lemma'),
                    help='分かち書きのトークン列（既定は語彙素 tokens_lemma。07 と同じ）')
    ap.add_argument('--meta', default=None, help='既定: *_v3_local.csv > v3 > v2')
    ap.add_argument('--questioned', default=DEFAULT_QUESTIONED,
                    help='作家を伏せる「問題のテクスト」（作家:題）')
    ap.add_argument('--known', default=DEFAULT_KNOWN,
                    help='既知の作品（作家:題 をカンマ区切り。各作家2点を推奨）')
    ap.add_argument('--mfw', type=int, default=50, help='最頻語の数')
    ap.add_argument('--label', default='作品X', help='問題のテクストの表示名')
    ap.add_argument('--out', default=os.path.join(ROOT, 'my_work', 'results',
                                                  'Step4_delta_manual.xlsx'))
    args = ap.parse_args()

    try:
        from openpyxl import Workbook
        from openpyxl.comments import Comment
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter as L
    except ImportError:
        sys.exit('openpyxl が必要である: uv pip install openpyxl')

    meta = args.meta or default_meta()
    rows = list(csv.DictReader(open(meta, encoding='utf-8-sig')))
    q = resolve(args.questioned, rows, args.tokens)
    known = [resolve(s, rows, args.tokens) for s in args.known.split(',') if s.strip()]
    if q['stem'] in {k['stem'] for k in known}:
        sys.exit('問題のテクストが既知の作品にも入っている')
    authors = list(dict.fromkeys(k['author'] for k in known))
    if len(known) < 3:
        sys.exit('既知の作品は3点以上にすること（標準偏差が意味を持たない）')

    def count(p):
        with open(p, encoding='utf-8') as fh:
            return Counter(fh.read().split())

    C = {k['stem']: count(k['path']) for k in known}
    Cq = count(q['path'])
    total = Counter()
    for c in C.values():
        total.update(c)
    mfw = [w for w, _ in total.most_common(args.mfw)]
    nk, nw = len(known), len(mfw)

    # ---- 正解値（ノートブックの照合用）-----------------------------------
    NT = np.array([sum(C[k['stem']].values()) for k in known], float)
    nq = float(sum(Cq.values()))
    F = np.array([[C[k['stem']][w] for w in mfw] for k in known]) / NT[:, None] * 1e4
    fq = np.array([Cq[w] for w in mfw]) / nq * 1e4
    mu, sd = F.mean(0), F.std(0)            # ddof=0（STDEVP と同じ）
    sd[sd == 0] = 1e-12
    Z, zq = (F - mu) / sd, (fq - mu) / sd
    delta_w = np.abs(Z - zq).mean(1)
    prof = {a: Z[[i for i, k in enumerate(known) if k['author'] == a]].mean(0)
            for a in authors}
    delta_a = {a: float(np.abs(v - zq).mean()) for a, v in prof.items()}

    # ---- ブック ------------------------------------------------------------
    FONT = 'Yu Gothic'
    f_in = Font(name=FONT, color='0000FF')          # 入力値（書き換えてよい）
    f_fx = Font(name=FONT, color='000000')          # 数式
    f_hd = Font(name=FONT, bold=True)
    f_ti = Font(name=FONT, bold=True, size=14)
    f_no = Font(name=FONT, color='555555', size=9)
    fill_hd = PatternFill('solid', fgColor='E8EEF4')
    fill_q = PatternFill('solid', fgColor='FFF2CC')  # 作品X の列
    fill_y = PatternFill('solid', fgColor='FFFF00')  # 触ってよいセル
    thin = Border(bottom=Side(style='thin', color='999999'))
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    wb = Workbook()
    ws0 = wb.active
    ws0.title = '説明'

    # 共通の配置：A=順位, B=語, C..=既知の作品, 最後の列=作品X
    R_AU, R_TI, R_N, R0 = 3, 4, 5, 6               # 作家・作品・総語数・語の開始行
    R1 = R0 + nw - 1
    cols = [L(3 + i) for i in range(nk)]            # 既知の作品の列
    CQ = L(3 + nk)                                  # 作品X の列
    labels = [f"{k['author']}『{k['title']}』" for k in known]

    def frame(ws, title, note, numfmt, with_n=True):
        ws['A1'] = title
        ws['A1'].font = f_ti
        ws['A2'] = note
        ws['A2'].font = f_no
        ws.cell(R_AU, 1, '作家').font = f_hd
        ws.cell(R_TI, 1, '作品').font = f_hd
        ws.cell(R_TI, 2, '語（語彙素）').font = f_hd
        for i, k in enumerate(known):
            ws.cell(R_AU, 3 + i, k['author'])
            ws.cell(R_TI, 3 + i, k['title'])
        ws.cell(R_AU, 3 + nk, '？')
        ws.cell(R_TI, 3 + nk, args.label)
        for r in (R_AU, R_TI):
            for c in range(1, 4 + nk):
                cell = ws.cell(r, c)
                cell.fill = fill_hd
                cell.alignment = center
                if c >= 3:
                    cell.font = f_hd
        ws.cell(R_AU, 3 + nk).fill = fill_q
        ws.cell(R_TI, 3 + nk).fill = fill_q
        if with_n:
            ws.cell(R_N, 1, '総語数').font = f_hd
        for j, w in enumerate(mfw):
            ws.cell(R0 + j, 1, j + 1).font = f_fx
            c = ws.cell(R0 + j, 2, disp(w))
            c.font = f_hd
            if disp(w) != w:
                c.comment = Comment(f'UniDic の語彙素：{w}', 'JLit')
        ws.column_dimensions['A'].width = 7
        ws.column_dimensions['B'].width = 12
        for c in range(3, 4 + nk):
            ws.column_dimensions[L(c)].width = 13
        ws.row_dimensions[R_TI].height = 32
        ws.freeze_panes = ws.cell(R0, 3)
        return ws

    # ---- 1_度数 -------------------------------------------------------------
    ws1 = frame(wb.create_sheet('1_度数'), '手順 0：最頻語の度数',
                f'既知の {nk} 作品を合わせた最頻語 {nw} 語。度数と総語数は入力値（青字）。'
                '語は語彙素なので「する」は「為る」，「いる」は「居る」と表示される。', '0')
    for i, k in enumerate(known):
        c = ws1.cell(R_N, 3 + i, int(NT[i]))
        c.font, c.number_format = f_in, '#,##0'
    c = ws1.cell(R_N, 3 + nk, int(nq))
    c.font, c.number_format = f_in, '#,##0'
    for j, w in enumerate(mfw):
        for i, k in enumerate(known):
            c = ws1.cell(R0 + j, 3 + i, C[k['stem']][w])
            c.font, c.number_format = f_in, '#,##0'
        c = ws1.cell(R0 + j, 3 + nk, Cq[w])
        c.font, c.number_format = f_in, '#,##0'
        ws1.cell(R0 + j, 3 + nk).fill = fill_q
    ws1.cell(R1 + 2, 1, '最頻語は既知の作品だけから選んだ。作品X は物差しを決めることに加わらない。').font = f_no

    # ---- 2_相対頻度 ---------------------------------------------------------
    ws2 = frame(wb.create_sheet('2_相対頻度'), '手順 1：1万語あたりの相対頻度',
                '式：度数/総語数×10000。作品の長さの違いを取り除く。', '0.00')
    for c_ in cols + [CQ]:
        ws2[f'{c_}{R_N}'] = f"='1_度数'!{c_}{R_N}"
        ws2[f'{c_}{R_N}'].number_format = '#,##0'
        ws2[f'{c_}{R_N}'].font = Font(name=FONT, color='008000')
        for r in range(R0, R1 + 1):
            ws2[f'{c_}{r}'] = f"='1_度数'!{c_}{r}/'1_度数'!{c_}${R_N}*10000"
            ws2[f'{c_}{r}'].number_format = '0.00'
            ws2[f'{c_}{r}'].font = f_fx
    for r in range(R0, R1 + 1):
        ws2[f'{CQ}{r}'].fill = fill_q

    # ---- 3_平均と標準偏差 ---------------------------------------------------
    ws3 = wb.create_sheet('3_平均と標準偏差')
    ws3['A1'] = '手順 2a：語ごとの平均と標準偏差（既知の作品だけ）'
    ws3['A1'].font = f_ti
    ws3['A2'] = (f'平均 =AVERAGE(既知 {nk} 作品)，標準偏差 =STDEVP(既知 {nk} 作品)。'
                 '作品X は含めない。変動係数＝標準偏差/平均 が大きい語ほど作品によって揺れる。')
    ws3['A2'].font = f_no
    hdr = ['順位', '語（語彙素）', '平均', '標準偏差', '変動係数', f'{args.label}の相対頻度']
    for c, h in enumerate(hdr, 1):
        cell = ws3.cell(R_TI, c, h)
        cell.font, cell.fill, cell.alignment = f_hd, fill_hd, center
    rng = lambda r: f"'2_相対頻度'!{cols[0]}{r}:{cols[-1]}{r}"
    for j, w in enumerate(mfw):
        r = R0 + j
        ws3.cell(r, 1, j + 1)
        c = ws3.cell(r, 2, disp(w))
        c.font = f_hd
        if disp(w) != w:
            c.comment = Comment(f'UniDic の語彙素：{w}', 'JLit')
        ws3[f'C{r}'] = f'=AVERAGE({rng(r)})'
        ws3[f'D{r}'] = f'=STDEVP({rng(r)})'
        ws3[f'E{r}'] = f'=IF(C{r}=0,0,D{r}/C{r})'
        ws3[f'F{r}'] = f"='2_相対頻度'!{CQ}{r}"
        for col, fmt in (('C', '0.00'), ('D', '0.00'), ('E', '0.000'), ('F', '0.00')):
            ws3[f'{col}{r}'].number_format = fmt
            ws3[f'{col}{r}'].font = f_fx
        ws3[f'F{r}'].font = Font(name=FONT, color='008000')
        ws3[f'F{r}'].fill = fill_q
    for col, wdt in zip('ABCDEF', (7, 12, 11, 11, 10, 16)):
        ws3.column_dimensions[col].width = wdt
    ws3.row_dimensions[R_TI].height = 32
    ws3.freeze_panes = 'C6'
    ws3.conditional_formatting.add(
        f'E{R0}:E{R1}', ColorScaleRule(start_type='min', start_color='FFFFFF',
                                       end_type='max', end_color='F4B183'))

    # ---- 4_zスコア ----------------------------------------------------------
    ws4 = frame(wb.create_sheet('4_zスコア'), '手順 2b：z スコア',
                '式：(相対頻度−平均)/標準偏差。頻度の高い「の」も低い語も，同じ尺度（標準偏差いくつ分）になる。'
                '赤＝平均より多い，青＝少ない。', '0.000', with_n=False)
    for c_ in cols + [CQ]:
        for r in range(R0, R1 + 1):
            ws4[f'{c_}{r}'] = (f"=('2_相対頻度'!{c_}{r}-'3_平均と標準偏差'!$C{r})"
                               f"/'3_平均と標準偏差'!$D{r}")
            ws4[f'{c_}{r}'].number_format = '0.000'
            ws4[f'{c_}{r}'].font = f_fx
    zscale = ColorScaleRule(start_type='num', start_value=-2, start_color='5B9BD5',
                            mid_type='num', mid_value=0, mid_color='FFFFFF',
                            end_type='num', end_value=2, end_color='E06666')
    ws4.conditional_formatting.add(f'{cols[0]}{R0}:{CQ}{R1}', zscale)
    ws4.cell(R_TI, 3 + nk).comment = Comment(
        '作品X の z は，既知の作品の平均と標準偏差で測っている。'
        '|z| が 2 を超える語は，既知の作品の範囲から外れている。', 'JLit')

    # ---- 5_z差 --------------------------------------------------------------
    # 左：作品ごと |z(X) − z(候補)|，右：作家プロファイル（2作品の z の平均）との差
    ws5 = frame(wb.create_sheet('5_z差'), '手順 3a：作品X との z の差（絶対値）',
                '式：ABS(z(作品X) − z(候補))。値が大きい語ほど，その候補と作品X を遠ざけている。'
                '右側は作家ごとのプロファイル（その作家の作品の z の平均）との差。', '0.000',
                with_n=False)
    ws5.cell(R_TI, 3 + nk, '（空欄）')
    ws5.cell(R_AU, 3 + nk, '')
    for c_ in cols:
        for r in range(R0, R1 + 1):
            ws5[f'{c_}{r}'] = f"=ABS('4_zスコア'!${CQ}{r}-'4_zスコア'!{c_}{r})"
            ws5[f'{c_}{r}'].number_format = '0.000'
            ws5[f'{c_}{r}'].font = f_fx
    A0 = 3 + nk + 2                                 # 作家プロファイルの開始列
    PZ = {}                                         # 作家 → プロファイル z の列
    PD = {}                                         # 作家 → |X − プロファイル| の列
    for ai, a in enumerate(authors):
        members = [cols[i] for i, k in enumerate(known) if k['author'] == a]
        cz, cd = L(A0 + 2 * ai), L(A0 + 2 * ai + 1)
        PZ[a], PD[a] = cz, cd
        for c_, t in ((cz, f'{a}\nプロファイルの z'), (cd, f'|X − {a}|')):
            ws5[f'{c_}{R_TI}'] = t
            ws5[f'{c_}{R_TI}'].font, ws5[f'{c_}{R_TI}'].fill = f_hd, fill_hd
            ws5[f'{c_}{R_TI}'].alignment = center
            ws5[f'{c_}{R_AU}'] = a
            ws5[f'{c_}{R_AU}'].font, ws5[f'{c_}{R_AU}'].fill = f_hd, fill_hd
            ws5[f'{c_}{R_AU}'].alignment = center
            ws5.column_dimensions[c_].width = 13
        for r in range(R0, R1 + 1):
            ws5[f'{cz}{r}'] = ('=AVERAGE(' + ','.join(f"'4_zスコア'!{m}{r}" for m in members) + ')')
            ws5[f'{cd}{r}'] = f"=ABS('4_zスコア'!${CQ}{r}-{cz}{r})"
            for c_ in (cz, cd):
                ws5[f'{c_}{r}'].number_format = '0.000'
                ws5[f'{c_}{r}'].font = f_fx
        ws5.conditional_formatting.add(f'{cd}{R0}:{cd}{R1}', ColorScaleRule(
            start_type='num', start_value=0, start_color='FFFFFF',
            end_type='num', end_value=3, end_color='E06666'))
    ws5.conditional_formatting.add(f'{cols[0]}{R0}:{cols[-1]}{R1}', ColorScaleRule(
        start_type='num', start_value=0, start_color='FFFFFF',
        end_type='num', end_value=3, end_color='E06666'))
    ws5.row_dimensions[R_TI].height = 44

    # ---- 6_Delta ------------------------------------------------------------
    ws6 = wb.create_sheet('6_Delta')
    ws6['A1'] = f'手順 3b：Delta ＝ z の差の絶対値の平均（{nw} 語）'
    ws6['A1'].font = f_ti
    ws6['A2'] = ('式：AVERAGE(5_z差 の列)。小さいほど作品X に近い。'
                 '「最も効いている語」は，その候補と作品X をいちばん遠ざけている語。')
    ws6['A2'].font = f_no
    h6 = ['候補', '作家', '作品', 'Delta', '順位（近い順）', '最も効いている語', 'その語の |z差|']
    for c, h in enumerate(h6, 1):
        cell = ws6.cell(4, c, h)
        cell.font, cell.fill, cell.alignment = f_hd, fill_hd, center
    words = f"'5_z差'!$B${R0}:$B${R1}"
    r6 = 5
    first_w = r6
    for i, k in enumerate(known):
        col = f"'5_z差'!{cols[i]}${R0}:{cols[i]}${R1}"
        ws6.cell(r6, 1, '作品')
        ws6.cell(r6, 2, k['author'])
        ws6.cell(r6, 3, k['title'])
        ws6[f'D{r6}'] = f'=AVERAGE({col})'
        ws6[f'E{r6}'] = f'=RANK(D{r6},$D${first_w}:$D${first_w + nk - 1},1)'
        ws6[f'F{r6}'] = f'=INDEX({words},MATCH(MAX({col}),{col},0))'
        ws6[f'G{r6}'] = f'=MAX({col})'
        r6 += 1
    last_w = r6 - 1
    r6 += 1
    first_a = r6
    for a in authors:
        col = f"'5_z差'!{PD[a]}${R0}:{PD[a]}${R1}"
        ws6.cell(r6, 1, '作家プロファイル')
        ws6.cell(r6, 2, a)
        ws6.cell(r6, 3, '（' + '・'.join(k['title'] for k in known if k['author'] == a) + ' の z の平均）')
        ws6[f'D{r6}'] = f'=AVERAGE({col})'
        ws6[f'E{r6}'] = f'=RANK(D{r6},$D${first_a}:$D${first_a + len(authors) - 1},1)'
        ws6[f'F{r6}'] = f'=INDEX({words},MATCH(MAX({col}),{col},0))'
        ws6[f'G{r6}'] = f'=MAX({col})'
        r6 += 1
    last_a = r6 - 1
    for r in list(range(first_w, last_w + 1)) + list(range(first_a, last_a + 1)):
        for c in 'ABCDEFG':
            ws6[f'{c}{r}'].font = f_fx if c in 'DEFG' else Font(name=FONT)
            ws6[f'{c}{r}'].border = thin
        ws6[f'D{r}'].number_format = '0.0000'
        ws6[f'G{r}'].number_format = '0.000'
        ws6[f'E{r}'].alignment = center
    for rr in ((first_w, last_w), (first_a, last_a)):
        ws6.conditional_formatting.add(f'D{rr[0]}:D{rr[1]}', ColorScaleRule(
            start_type='min', start_color='63BE7B', end_type='max', end_color='FFFFFF'))
    for col, wdt in zip('ABCDEFG', (18, 12, 34, 10, 12, 16, 12)):
        ws6.column_dimensions[col].width = wdt
    ws6.cell(last_a + 2, 1,
             '問い：作品X にいちばん近い作品と，いちばん近い作家プロファイルは一致するか。'
             '一致しないなら，それはなぜか（5_z差 で色の濃い語を見てみよう）。').font = f_no

    # ---- 7_語数Nを変える -----------------------------------------------------
    ws7 = wb.create_sheet('7_語数Nを変える')
    ws7['A1'] = '手順 4：使う語の数 N を変える'
    ws7['A1'].font = f_ti
    ws7['A2'] = (f'黄色のセルに 1〜{nw} の数を入れる。上位 N 語だけで Delta を計算し直す。'
                 '式：SUMPRODUCT((順位<=N)*|z差|)/N。N が小さいと順位はどう揺れるか。')
    ws7['A2'].font = f_no
    ws7['A3'] = 'N ='
    ws7['A3'].font = f_hd
    ws7['B3'] = min(20, nw)
    ws7['B3'].font, ws7['B3'].fill = f_in, fill_y
    ws7['C3'] = f'（1〜{nw}）'
    ws7['C3'].font = f_no
    ranks = f"'5_z差'!$A${R0}:$A${R1}"
    Ns = [n for n in (10, 20, 30, 40, 50, 75, 100, 150, 200, 300) if n <= nw]
    if nw not in Ns:
        Ns.append(nw)
    h7 = ['候補', '作家', '作品', 'Delta（N 語）', '順位'] + [f'N={n}' for n in Ns]
    for c, h in enumerate(h7, 1):
        cell = ws7.cell(5, c, h)
        cell.font, cell.fill, cell.alignment = f_hd, fill_hd, center
    cands = ([('作品', k['author'], k['title'], cols[i]) for i, k in enumerate(known)]
             + [('作家プロファイル', a, '（平均）', PD[a]) for a in authors])
    r7 = 6
    groups = [(r7, r7 + nk - 1), (r7 + nk + 1, r7 + nk + len(authors))]
    for gi, (g0, g1) in enumerate(groups):
        sub = cands[:nk] if gi == 0 else cands[nk:]
        for (kind, a, t, colL), r in zip(sub, range(g0, g1 + 1)):
            vals = f"'5_z差'!{colL}${R0}:{colL}${R1}"
            ws7.cell(r, 1, kind)
            ws7.cell(r, 2, a)
            ws7.cell(r, 3, t)
            ws7[f'D{r}'] = f'=SUMPRODUCT(({ranks}<=$B$3)*{vals})/$B$3'
            ws7[f'E{r}'] = f'=RANK(D{r},$D${g0}:$D${g1},1)'
            ws7[f'D{r}'].number_format = '0.0000'
            ws7[f'E{r}'].alignment = center
            for ci, n in enumerate(Ns):
                cl = L(6 + ci)
                ws7[f'{cl}{r}'] = f'=SUMPRODUCT(({ranks}<={n})*{vals})/{n}'
                ws7[f'{cl}{r}'].number_format = '0.000'
                ws7[f'{cl}{r}'].font = f_fx
            for c in range(1, 6 + len(Ns)):
                ws7.cell(r, c).border = thin
        for ci in range(len(Ns)):
            cl = L(6 + ci)
            ws7.conditional_formatting.add(f'{cl}{g0}:{cl}{g1}', ColorScaleRule(
                start_type='min', start_color='63BE7B', end_type='max', end_color='FFFFFF'))
    for col, wdt in zip('ABCDE', (18, 12, 22, 13, 8)):
        ws7.column_dimensions[col].width = wdt
    for ci in range(len(Ns)):
        ws7.column_dimensions[L(6 + ci)].width = 9
    ws7.cell(groups[1][1] + 2, 1,
             '右の N=… の列は固定の N。緑が濃いほど近い。色の並びが N によって入れ替わるかを見る。').font = f_no

    # ---- 説明 --------------------------------------------------------------
    lines = [
        ("Burrows's Delta を手で計算する（Step 4 §2）", f_ti),
        ('', None),
        (f'問題のテクスト：{args.label}（作家は伏せてある）', f_hd),
        ('既知の作品：' + '，'.join(labels), None),
        (f'最頻語：既知の {nk} 作品を合わせた上位 {nw} 語（語彙素）', None),
        ('', None),
        ('手順とシート', f_hd),
        ('  1_度数          最頻語の度数と総語数（入力値）', None),
        ('  2_相対頻度      1 万語あたりにそろえる', None),
        ('  3_平均と標準偏差 語ごとに，既知の作品だけから平均と標準偏差を出す', None),
        ('  4_zスコア       (相対頻度−平均)/標準偏差。語の頻度の大小を消し，どの語も同じ重みにする', None),
        ('  5_z差           作品X と各候補の z の差の絶対値（語ごと）', None),
        ('  6_Delta         5 の平均＝Delta。小さいほど近い', None),
        ('  7_語数Nを変える  上位 N 語だけで計算し直す（黄色のセルに N を入れる）', None),
        ('', None),
        ('凡例', f_hd),
        ('  青字＝入力値（書き換えてよい。書き換えると全シートが再計算される）', Font(name=FONT, color='0000FF')),
        ('  黒字＝数式／緑字＝別シートの値をそのまま引いたもの', None),
        ('  黄色のセル＝ここに値を入れて試す（7_語数Nを変える の N）', None),
        ('  薄い黄色の列＝作品X', None),
        ('', None),
        ('約束', f_hd),
        ('  ・最頻語・平均・標準偏差は既知の作品だけから決める。作品X が物差しを動かさないため（Burrows 2002）', None),
        ('  ・標準偏差は母標準偏差（STDEVP）。ノートブックの numpy（ddof=0）と 07_descriptive_stats.py に揃えてある', None),
        ('  ・並べ替え（ソート）はしないこと。数式の参照が崩れる。並べて見たいときは別のシートに値を貼り付ける', None),
        ('', None),
        ('考えること', f_hd),
        ('  1. 6_Delta で作品X にいちばん近い作品と作家プロファイルはどれか。両者は一致するか', None),
        ('  2. 5_z差 で色の濃い語は何か。それは作家の癖か，語りの人称や登場人物名など別の要因か', None),
        ('  3. 7_語数Nを変える で N を 10・20・50 と変えると，順位はどう変わるか', None),
        ('  4. 1_度数 の数字を1つ書き換えて，どこまで影響が伝わるかを追ってみる', None),
        ('', None),
        (f'作成：17_delta_workbook.py（メタデータ {os.path.basename(meta)}，トークン {os.path.basename(args.tokens.rstrip("/"))}）', f_no),
        ('参考：Burrows, J. (2002) “Delta”: a measure of stylistic difference and a guide to likely authorship. LLC 17(3): 267–287.', f_no),
    ]
    for i, (t, f) in enumerate(lines, 1):
        ws0.cell(i, 1, t).font = f or Font(name=FONT)
    ws0.column_dimensions['A'].width = 110

    os.makedirs(os.path.dirname(os.path.abspath(args.out)) or '.', exist_ok=True)
    wb.save(args.out)

    exp = {
        'questioned': {'label': args.label, **{k: q[k] for k in ('author', 'title', 'stem')}},
        'known': [{k2: k[k2] for k2 in ('author', 'title', 'stem')} for k in known],
        'mfw': mfw, 'ddof': 0, 'per': 10000,
        'delta_work': {f"{k['author']}『{k['title']}』": float(d) for k, d in zip(known, delta_w)},
        'delta_author': delta_a,
    }
    jpath = os.path.splitext(args.out)[0] + '_expected.json'
    with open(jpath, 'w', encoding='utf-8') as fh:
        json.dump(exp, fh, ensure_ascii=False, indent=1)

    print(f'[ok  ] {args.out}')
    print(f'       既知 {nk} 作品（{len(authors)} 作家）・最頻語 {nw} 語・問題のテクスト＝{args.label}')
    print(f'       照合用の正解値 → {jpath}')
    incomplete = [f"{k['author']}『{k['title']}』" for k in known + [q]
                  if k['complete'] and k['complete'] != 'complete']
    if incomplete:
        print('[warn] 完本でない作品を含む: ' + '，'.join(incomplete))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
